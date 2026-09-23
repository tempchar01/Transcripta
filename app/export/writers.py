from __future__ import annotations

from datetime import datetime
from pathlib import Path

from app.core.models import Transcript
from app.transcription.readable import ReadableTranscript, build_readable_transcript


def _readable(transcript: Transcript, readable: ReadableTranscript | None) -> ReadableTranscript:
    return readable or build_readable_transcript(transcript)


def write_docx(transcript: Transcript, target: Path, readable: ReadableTranscript | None = None) -> Path:
    try:
        from docx import Document
    except ImportError as error:
        raise RuntimeError("Для DOCX установите python-docx.") from error
    target.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("Transcripta", level=0)
    document.add_paragraph(f"Исходный файл: {Path(transcript.source_file).name}")
    document.add_paragraph(f"Дата: {datetime.fromisoformat(transcript.created_at).astimezone().strftime('%d.%m.%Y %H:%M')}")
    presentation = _readable(transcript, readable)
    previous_speaker: str | None | object = object()
    for paragraph in presentation.paragraphs:
        if paragraph.speaker_id is not None and paragraph.speaker_id != previous_speaker:
            document.add_heading(presentation.speaker_label(paragraph.speaker_id), level=2)
        document.add_paragraph(paragraph.text)
        previous_speaker = paragraph.speaker_id
    document.save(target)
    return target


def export_docx(transcript: Transcript, output_dir: Path, readable: ReadableTranscript | None = None) -> dict[str, Path]:
    stem = Path(transcript.source_file).stem
    output_dir.mkdir(parents=True, exist_ok=True)
    sequence = 1
    while True:
        basename = stem if sequence == 1 else f"{stem} ({sequence})"
        try:
            target = output_dir / f"{basename}.docx"
            with target.open("xb"):
                pass
            break
        except FileExistsError:
            sequence += 1
    try:
        return {"docx": write_docx(transcript, target, readable)}
    except Exception:
        target.unlink(missing_ok=True)
        raise
    return result
