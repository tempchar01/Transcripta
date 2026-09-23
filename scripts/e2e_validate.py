from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import asdict
from pathlib import Path

from app.audio.media import ffmpeg_resolution, get_duration
from app.config.settings import AppSettings
from app.jobs.checkpoint import CheckpointStore
from app.jobs.runner import JobCancelled, Progress, TranscriptionJob


def _validate_exports(paths: dict[str, Path]) -> dict[str, str | bool]:
    result: dict[str, str | bool] = {}
    for name, path in paths.items():
        result[f"{name}_exists"] = path.is_file() and path.stat().st_size > 0
    if "txt" in paths:
        result["txt_utf8"] = bool(paths["txt"].read_text(encoding="utf-8").strip())
    if "json" in paths:
        data = json.loads(paths["json"].read_text(encoding="utf-8"))
        result["json_has_segments"] = bool(data.get("segments"))
    if "srt" in paths:
        content = paths["srt"].read_text(encoding="utf-8")
        timings = re.findall(r"(\d\d:\d\d:\d\d,\d{3}) --> (\d\d:\d\d:\d\d,\d{3})", content)
        result["srt_valid_intervals"] = bool(timings) and all(start <= end for start, end in timings)
    if "docx" in paths:
        from docx import Document
        result["docx_has_text"] = bool("\n".join(item.text for item in Document(paths["docx"]).paragraphs).strip())
    return result


def _export_segment_counts(paths: dict[str, Path]) -> dict[str, int | None]:
    counts: dict[str, int | None] = {"txt": None, "docx": None, "json": None, "srt": None}
    if "txt" in paths:
        counts["txt"] = len(re.findall(r"^\[\d\d:\d\d:\d\d\]$", paths["txt"].read_text(encoding="utf-8"), re.MULTILINE))
    if "json" in paths:
        counts["json"] = len(json.loads(paths["json"].read_text(encoding="utf-8")).get("segments", []))
    if "srt" in paths:
        counts["srt"] = len(re.findall(r"^\d+$", paths["srt"].read_text(encoding="utf-8"), re.MULTILINE))
    if "docx" in paths:
        from docx import Document
        counts["docx"] = len(re.findall(r"^\[\d\d:\d\d:\d\d\]$", "\n".join(item.text for item in Document(paths["docx"]).paragraphs), re.MULTILINE))
    return counts


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="Run the real Transcripta PyAV → chunk → ASR → checkpoint → export pipeline.")
    parser.add_argument("audio", type=Path)
    parser.add_argument("--app-data", type=Path, default=Path("work") / "e2e-app-data")
    parser.add_argument("--output", type=Path, default=Path("outputs") / "e2e")
    parser.add_argument("--chunk-seconds", type=int, default=45, help="Internal integration override; production AUTO policy remains unchanged.")
    parser.add_argument("--model", choices=("auto", "medium", "large"), default="auto")
    parser.add_argument("--profile", choices=("auto", "memory_saver", "balanced", "quality", "cpu"), default="auto")
    parser.add_argument("--allow-model-download", action="store_true")
    parser.add_argument("--no-vad", action="store_true", help="Diagnostic control: disable VAD for this run.")
    parser.add_argument("--cancel-after-chunks", type=int, help="Safely cancel after this many completed checkpoints.")
    args = parser.parse_args()
    try:
        duration = get_duration(args.audio)
    except Exception as error:
        print(json.dumps({"status": "blocked", "media_backend": "pyav", "ffmpeg": ffmpeg_resolution(), "error": str(error)}, ensure_ascii=False, indent=2))
        return 1
    settings = AppSettings()
    settings.recognition.model = args.model
    settings.recognition.performance_profile = args.profile
    settings.recognition.chunk_seconds = args.chunk_seconds
    settings.recognition.vad_filter = not args.no_vad
    settings.export.output_dir = str(args.output)
    job: TranscriptionJob | None = None
    progress_events: list[Progress] = []

    def on_progress(progress: Progress) -> None:
        progress_events.append(progress)
        if args.cancel_after_chunks and len(progress_events) >= args.cancel_after_chunks and job is not None:
            job.cancel()

    job = TranscriptionJob(args.audio, duration, settings, args.app_data, args.allow_model_download, on_progress=on_progress,
                           on_warning=lambda message: print(f"WARNING: {message}", file=sys.stderr))
    report: dict[str, object] = {"media_backend": "pyav", "ffmpeg": ffmpeg_resolution(), "audio_duration_seconds": duration}
    try:
        result = job.run()
    except JobCancelled:
        checkpoint = CheckpointStore(job._job_directory() / "checkpoint.json").load()
        report.update({"status": "cancelled", "completed_chunks": len(progress_events),
                       "checkpoint": checkpoint.to_dict() if checkpoint else None})
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 2
    except Exception as error:
        report.update({"status": "blocked", "stage": "transcription", "error": str(error)})
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 1
    report.update({
        "status": "passed", "model": result.transcript.model, "compute_type": result.hardware_profile.compute_type,
        "model_id": job.model_manager.model_id(result.transcript.model),
        "gpu": result.device, "gpu_inference_succeeded": result.gpu_inference_succeeded,
        "processing_seconds": result.processing_seconds,
        "realtime_factor": result.processing_seconds / duration if duration else None,
        "peak_ram_mb": result.peak_ram_mb, "peak_vram_mb_sampled": result.peak_vram_mb,
        "chunks": len(progress_events), "segments": len(result.transcript.segments),
        "transcript_characters": sum(len(item.text) for item in result.transcript.segments),
        "decode": asdict(result.decode_diagnostics), "processed_coverage_seconds": result.coverage.covered_duration,
        "coverage_gaps": result.coverage.gaps, "completed_chunks": [asdict(chunk) for chunk in result.completed_chunks],
        "first_timestamp_seconds": result.transcript.segments[0].start if result.transcript.segments else None,
        "last_timestamp_seconds": result.transcript.segments[-1].end if result.transcript.segments else None,
        "exports": {name: str(path) for name, path in result.exported.items()},
        "export_validation": _validate_exports(result.exported), "export_segment_counts": _export_segment_counts(result.exported),
    })
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
