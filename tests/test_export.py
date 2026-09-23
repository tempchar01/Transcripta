from docx import Document

from app.core.models import Segment, Transcript
from app.export.writers import export_docx, write_docx


def sample_transcript() -> Transcript:
    return Transcript("meeting.mp3", 30.0, "ru", "small", created_at="2026-09-06T10:00:00+00:00", segments=[
        Segment(1, 4.12, 8.43, "Сегодня мы обсудим."),
        Segment(2, 8.43, 13.12, "Следующий вопрос."),
        Segment(3, 14.0, 14.0, ""),
    ])


def test_repeated_unicode_docx_export_preserves_previous_result(tmp_path):
    transcript = sample_transcript()
    transcript.source_file = "Запись встречи.ogg"
    first = export_docx(transcript, tmp_path)
    original = first["docx"].read_bytes()
    second = export_docx(transcript, tmp_path)
    assert second["docx"].stem == "Запись встречи (2)"
    assert first["docx"].read_bytes() == original
    assert {path.suffix for path in tmp_path.iterdir()} == {".docx"}


def test_docx_export(tmp_path):
    target = write_docx(sample_transcript(), tmp_path / "result.docx")
    document = Document(target)
    text = "\n".join(item.text for item in document.paragraphs)
    assert "Transcripta" in text
    assert "Сегодня мы обсудим." in text
    assert "[00:00:04]" not in text


def test_docx_preserves_one_speaker_heading_for_one_turn(tmp_path):
    transcript = sample_transcript()
    transcript.segments[0].speaker_id = transcript.segments[1].speaker_id = "1"
    target = write_docx(transcript, tmp_path / "speakers.docx")
    text = "\n".join(item.text for item in Document(target).paragraphs)
    assert text.count("Спикер 1") == 1
